import os
import logging
from datetime import datetime
from decimal import Decimal
from typing import Optional


from flask import current_app

from ..extensions import db
from ..models.project import Project

logger = logging.getLogger(__name__)


class ReportService:

    def _get_report_folder(self) -> str:
        folder = current_app.config.get('REPORT_FOLDER', 'reports')
        os.makedirs(folder, exist_ok=True)
        return folder

    def _default_formula_params(self, formula_params: dict = None) -> dict:
        defaults = {
            'loss_rate': 0.05,
            'tax_rate': 0.13,
            'aux_material_fee': 0,
            'labor_fee': 0,
        }
        if formula_params:
            defaults.update(formula_params)
        return defaults

    # ── Excel 核价单 ──────────────────────────────────────────────────────────

    def generate_internal_pricing_sheet(self, project_id: int, formula_params: dict = None) -> str:
        """生成 Excel 核价单，返回文件路径。"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        params = self._default_formula_params(formula_params)
        project = Project.query.get_or_404(project_id)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '核价单'

        # ── 样式定义 ──
        bold_font = Font(bold=True)
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(fill_type='solid', fgColor='366092')
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        def _apply_border(cell):
            cell.border = thin_border

        # ── 第1行：项目信息 ──
        ws.merge_cells('A1:H1')
        title_cell = ws['A1']
        project_date_str = project.project_date.isoformat() if project.project_date else ''
        title_cell.value = (
            f"项目：{project.name}    客户：{project.customer or ''}    日期：{project_date_str}"
        )
        title_cell.font = Font(bold=True, size=13)
        title_cell.alignment = center_align

        # ── 第2行：公式参数 ──
        ws.merge_cells('A2:H2')
        param_cell = ws['A2']
        param_cell.value = (
            f"折扣率：-    损耗率：{params['loss_rate']*100:.1f}%    "
            f"辅材费：{params['aux_material_fee']}    "
            f"人工费：{params['labor_fee']}    "
            f"税率：{params['tax_rate']*100:.1f}%"
        )
        param_cell.font = Font(italic=True, size=10)
        param_cell.alignment = Alignment(horizontal='left', vertical='center')

        # ── 第3行：表头 ──
        headers = ['层级', '型号', '名称', '规格', '数量', '单价', '折扣率', '总价']
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            _apply_border(cell)

        # ── 数据行 ──
        row_num = 4
        project_total = Decimal('0')

        for cabinet in project.cabinets:
            cabinet_total = Decimal('0')

            # Cabinet 行（加粗）
            cab_row_start = row_num
            ws.cell(row=row_num, column=1, value='配电柜').font = bold_font
            ws.cell(row=row_num, column=1).value = '配电柜'
            ws.cell(row=row_num, column=3, value=cabinet.name).font = bold_font
            for c in range(1, 9):
                cell = ws.cell(row=row_num, column=c)
                cell.font = bold_font
                _apply_border(cell)
            row_num += 1

            for sc in cabinet.structure_components:
                sc_total = Decimal('0')

                # StructureComponent 行（缩进1级）
                ws.cell(row=row_num, column=1, value='  结构件')
                ws.cell(row=row_num, column=3, value=sc.name)
                for c in range(1, 9):
                    _apply_border(ws.cell(row=row_num, column=c))
                row_num += 1

                for bc in sc.base_components:
                    bc_total = (
                        (bc.quantity or Decimal('0'))
                        * (bc.unit_price or Decimal('0'))
                        * (bc.discount_rate or Decimal('1'))
                    )
                    sc_total += bc_total

                    # BaseComponent 行（缩进2级）
                    ws.cell(row=row_num, column=1, value='    元器件')
                    ws.cell(row=row_num, column=2, value=bc.model_number or '')
                    ws.cell(row=row_num, column=3, value=bc.name or '')
                    ws.cell(row=row_num, column=4, value=bc.specification or '')
                    ws.cell(row=row_num, column=5, value=float(bc.quantity or 0))
                    ws.cell(row=row_num, column=6, value=float(bc.unit_price or 0))
                    ws.cell(row=row_num, column=7, value=float(bc.discount_rate or 1))
                    ws.cell(row=row_num, column=8, value=float(bc_total))
                    for c in range(1, 9):
                        _apply_border(ws.cell(row=row_num, column=c))
                    row_num += 1

                cabinet_total += sc_total

            # 配电柜小计行
            ws.cell(row=row_num, column=1, value='小计')
            ws.cell(row=row_num, column=3, value=cabinet.name)
            ws.cell(row=row_num, column=8, value=float(cabinet_total))
            for c in range(1, 9):
                cell = ws.cell(row=row_num, column=c)
                cell.font = bold_font
                _apply_border(cell)
            row_num += 1

            project_total += cabinet_total

        # ── 应用公式参数计算最终总价 ──
        loss_rate = Decimal(str(params['loss_rate']))
        tax_rate = Decimal(str(params['tax_rate']))
        aux_material_fee = Decimal(str(params['aux_material_fee']))
        labor_fee = Decimal(str(params['labor_fee']))
        final_total = project_total * (1 + loss_rate) * (1 + tax_rate) + aux_material_fee + labor_fee

        # ── 最后行：项目总价 ──
        ws.merge_cells(f'A{row_num}:G{row_num}')
        total_label = ws.cell(row=row_num, column=1, value='项目总价（含损耗、税费）')
        total_label.font = Font(bold=True, size=12)
        total_label.alignment = center_align
        total_value = ws.cell(row=row_num, column=8, value=float(final_total))
        total_value.font = Font(bold=True, size=12)
        for c in range(1, 9):
            _apply_border(ws.cell(row=row_num, column=c))

        # ── 列宽 ──
        col_widths = [12, 20, 25, 30, 10, 12, 10, 15]
        for i, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # ── 保存 ──
        timestamp = datetime.utcnow().strftime('%Y%m%d%H%M%S')
        filename = f'pricing_{project_id}_{timestamp}.xlsx'
        filepath = os.path.join(self._get_report_folder(), filename)
        wb.save(filepath)
        logger.info("Excel 核价单已生成: %s", filepath)
        return filepath

    # ── PDF 客户报价单 ────────────────────────────────────────────────────────

    def generate_customer_quotation(self, project_id: int, formula_params: dict = None) -> str:
        """生成 PDF 客户报价单，返回文件路径。"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfgen import canvas as rl_canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        params = self._default_formula_params(formula_params)
        project = Project.query.get_or_404(project_id)

        timestamp = datetime.utcnow().strftime('%Y%m%d%H%M%S')
        filename = f'quotation_{project_id}_{timestamp}.pdf'
        filepath = os.path.join(self._get_report_folder(), filename)

        company_name = current_app.config.get('COMPANY_NAME', '电气设备制造公司')

        # ── 尝试注册中文字体 ──
        _register_chinese_font()

        # ── 使用 canvas 直接绘制（支持水印） ──
        page_width, page_height = A4

        from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame

        class WatermarkDocTemplate(BaseDocTemplate):
            def __init__(self, filename, company, **kwargs):
                self.company = company
                super().__init__(filename, **kwargs)

            def handle_pageBegin(self):
                super().handle_pageBegin()
                self._draw_watermark_and_header()

            def _draw_watermark_and_header(self):
                c = self.canv
                # 水印
                c.saveState()
                c.setFont('Helvetica', 40)
                c.setFillColorRGB(0.85, 0.85, 0.85)
                c.setFillAlpha(0.3)
                c.translate(page_width / 2, page_height / 2)
                c.rotate(45)
                c.drawCentredString(0, 0, 'CONFIDENTIAL')
                c.restoreState()
                # 页眉
                c.saveState()
                font_name = _get_font_name()
                c.setFont(font_name, 10)
                c.setFillColorRGB(0.4, 0.4, 0.4)
                c.drawCentredString(page_width / 2, page_height - 1.5 * cm, self.company)
                c.setStrokeColorRGB(0.7, 0.7, 0.7)
                c.line(2 * cm, page_height - 1.8 * cm, page_width - 2 * cm, page_height - 1.8 * cm)
                c.restoreState()

        doc = WatermarkDocTemplate(
            filepath,
            company=company_name,
            pagesize=A4,
            leftMargin=2 * cm,
            rightMargin=2 * cm,
            topMargin=3 * cm,
            bottomMargin=2 * cm,
        )

        frame = Frame(
            doc.leftMargin, doc.bottomMargin,
            page_width - doc.leftMargin - doc.rightMargin,
            page_height - doc.topMargin - doc.bottomMargin,
            id='main',
        )
        doc.addPageTemplates([PageTemplate(id='main', frames=frame, onPage=_draw_page_decorations(company_name))])

        styles = getSampleStyleSheet()
        font_name = _get_font_name()

        title_style = ParagraphStyle(
            'CustomTitle', parent=styles['Title'],
            fontName=font_name, fontSize=16, spaceAfter=12,
        )
        normal_style = ParagraphStyle(
            'CustomNormal', parent=styles['Normal'],
            fontName=font_name, fontSize=10, spaceAfter=6,
        )
        heading_style = ParagraphStyle(
            'CustomHeading', parent=styles['Heading2'],
            fontName=font_name, fontSize=12, spaceAfter=8,
        )

        story = []

        # 标题
        story.append(Paragraph('客户报价单', title_style))
        story.append(Spacer(1, 0.3 * cm))

        # 项目信息
        project_date_str = project.project_date.isoformat() if project.project_date else '-'
        story.append(Paragraph(f'项目名称：{project.name}', normal_style))
        story.append(Paragraph(f'客户：{project.customer or "-"}', normal_style))
        story.append(Paragraph(f'日期：{project_date_str}', normal_style))
        story.append(Spacer(1, 0.5 * cm))

        # 配电柜汇总表
        story.append(Paragraph('配电柜报价汇总', heading_style))

        loss_rate = Decimal(str(params['loss_rate']))
        tax_rate = Decimal(str(params['tax_rate']))
        aux_material_fee = Decimal(str(params['aux_material_fee']))
        labor_fee = Decimal(str(params['labor_fee']))

        table_data = [['柜名', '报价（元）']]
        project_total = Decimal('0')

        for cabinet in project.cabinets:
            cabinet_total = Decimal('0')
            for sc in cabinet.structure_components:
                for bc in sc.base_components:
                    bc_total = (
                        (bc.quantity or Decimal('0'))
                        * (bc.unit_price or Decimal('0'))
                        * (bc.discount_rate or Decimal('1'))
                    )
                    cabinet_total += bc_total
            # 对每个柜应用公式
            cab_final = cabinet_total * (1 + loss_rate) * (1 + tax_rate)
            project_total += cab_final
            table_data.append([cabinet.name, f'{float(cab_final):,.2f}'])

        # 辅材费和人工费加到总价
        project_total += aux_material_fee + labor_fee

        table_data.append(['合计', f'{float(project_total):,.2f}'])

        col_widths_pdf = [(page_width - 4 * cm) * 0.6, (page_width - 4 * cm) * 0.4]
        table = Table(table_data, colWidths=col_widths_pdf)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('ALIGN', (0, 0), (0, -1), 'LEFT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#EEF2F7')]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#D9E1F2')),
            ('FONTNAME', (0, -1), (-1, -1), font_name),
            ('FONTSIZE', (0, -1), (-1, -1), 11),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.5 * cm))
        story.append(Paragraph(f'项目总价：{float(project_total):,.2f} 元', heading_style))

        doc.build(story)
        logger.info("PDF 客户报价单已生成: %s", filepath)
        return filepath

    # ── 文件管理 ──────────────────────────────────────────────────────────────

    def get_report_path(self, token: str)-> Optional[str]:
        """根据 token（文件名）获取报表文件路径。"""
        folder = current_app.config.get('REPORT_FOLDER', 'reports')
        filepath = os.path.join(folder, token)
        if os.path.isfile(filepath):
            return filepath
        return None


# ── 辅助函数 ──────────────────────────────────────────────────────────────────

_chinese_font_registered = False
_font_name = 'Helvetica'


def _register_chinese_font():
    """尝试注册系统中文字体，失败则回退到 Helvetica。"""
    global _chinese_font_registered, _font_name
    if _chinese_font_registered:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        # Linux
        '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
        '/usr/share/fonts/truetype/arphic/uming.ttc',
        # macOS
        '/System/Library/Fonts/PingFang.ttc',
        '/Library/Fonts/Arial Unicode MS.ttf',
        # Windows
        'C:/Windows/Fonts/msyh.ttc',
        'C:/Windows/Fonts/simhei.ttf',
        'C:/Windows/Fonts/simsun.ttc',
    ]
    for path in candidates:
        if os.path.isfile(path):
            try:
                pdfmetrics.registerFont(TTFont('ChineseFont', path))
                _font_name = 'ChineseFont'
                _chinese_font_registered = True
                logger.info("已注册中文字体: %s", path)
                return
            except Exception as exc:
                logger.warning("注册字体失败 %s: %s", path, exc)

    _chinese_font_registered = True  # 标记已尝试，避免重复
    logger.info("未找到中文字体，使用 Helvetica（中文可能显示为方块）")


def _get_font_name() -> str:
    return _font_name


def _draw_page_decorations(company_name: str):
    """返回 onPage 回调，用于绘制水印和页眉。"""
    from reportlab.lib.units import cm
    from reportlab.lib.pagesizes import A4

    page_width, page_height = A4

    def on_page(canvas, doc):
        # 水印
        canvas.saveState()
        canvas.setFont('Helvetica', 40)
        canvas.setFillColorRGB(0.85, 0.85, 0.85)
        try:
            canvas.setFillAlpha(0.3)
        except Exception:
            pass
        canvas.translate(page_width / 2, page_height / 2)
        canvas.rotate(45)
        canvas.drawCentredString(0, 0, 'CONFIDENTIAL')
        canvas.restoreState()

        # 页眉
        canvas.saveState()
        font_name = _get_font_name()
        canvas.setFont(font_name, 10)
        canvas.setFillColorRGB(0.4, 0.4, 0.4)
        canvas.drawCentredString(page_width / 2, page_height - 1.5 * cm, company_name)
        canvas.setStrokeColorRGB(0.7, 0.7, 0.7)
        canvas.line(2 * cm, page_height - 1.8 * cm, page_width - 2 * cm, page_height - 1.8 * cm)
        canvas.restoreState()

    return on_page


report_service = ReportService()
